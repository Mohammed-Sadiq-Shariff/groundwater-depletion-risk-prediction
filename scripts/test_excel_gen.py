import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import pandas as pd

def create_template_workbook():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Groundwater_Import_Template"

    # Header style
    header_fill = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    align_center = Alignment(horizontal="center", vertical="center")

    headers = [
        "Station", "District", "Tehsil", "Block", "Village",
        "Latitude", "Longitude", "Monitoring Date",
        "Groundwater Level (m)", "Recent Rainfall 30d (mm)"
    ]

    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center

    sample_rows = [
        ["Jayapura", "Mysore", "MYSORE", "MYSORE", "Jayapura", 12.2150, 76.5420, "2024-05-15", 8.42, 45.2],
        ["Hunsur Rural", "Mysore", "HUNSUR", "HUNSUR", "Hunsur", 12.3120, 76.2950, "2024-05-15", 6.18, 52.0],
        ["Nanjangud North", "Mysore", "NANJANGUD", "NANJANGUD", "Nanjangud", 12.1180, 76.6810, "2024-05-15", 7.05, 38.5],
        ["Bannur", "Mysore", "TIRUMAKUDAL-NARSIPUR", "TIRUMAKUDAL-NARSIPUR", "Bannur", 12.3167, 76.8500, "2024-05-15", 4.50, 61.0],
        ["Saligrama", "Mysore", "KRISHNARAJANAGARA", "KRISHNARAJANAGARA", "Saligrama", 12.5400, 76.2800, "2024-05-15", 9.10, 29.8]
    ]

    for row in sample_rows:
        ws.append(row)

    # Style data rows
    data_font = Font(name="Segoe UI", size=10)
    for row in ws.iter_rows(min_row=2, max_row=len(sample_rows) + 1, min_col=1, max_col=len(headers)):
        for cell in row:
            cell.font = data_font

    # Instructions sheet
    ws_info = wb.create_sheet(title="Field_Guidelines")
    ws_info.append(["Field Name", "Type", "Required", "Example", "Description"])
    info_header = ws_info[1]
    for cell in info_header:
        cell.fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        cell.font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")

    guidelines = [
        ["Station", "Text", "Yes", "Jayapura", "Unique name of the groundwater monitoring well"],
        ["District", "Text", "Yes", "Mysore", "District name (must be Mysore / Mysuru)"],
        ["Tehsil", "Text", "Yes", "MYSORE", "Taluk/Tehsil name"],
        ["Block", "Text", "Yes", "MYSORE", "Administrative block name"],
        ["Village", "Text", "Optional", "Jayapura", "Village location of the monitoring station"],
        ["Latitude", "Decimal", "Yes", "12.2150", "GPS Latitude (Decimal Degrees, 11.5 to 13.0)"],
        ["Longitude", "Decimal", "Yes", "76.5420", "GPS Longitude (Decimal Degrees, 75.5 to 77.5)"],
        ["Monitoring Date", "Date", "Yes", "2024-05-15", "Date of observation (YYYY-MM-DD)"],
        ["Groundwater Level (m)", "Decimal", "Yes", "8.42", "Depth to water level below ground surface in meters (mbgl)"],
        ["Recent Rainfall 30d (mm)", "Decimal", "Optional", "45.2", "Total rainfall in the preceding 30 days in millimeters"]
    ]

    for item in guidelines:
        ws_info.append(item)

    for sheet in [ws, ws_info]:
        for col in sheet.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            sheet.column_dimensions[col_letter].width = max(max_len + 3, 12)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

content = create_template_workbook()
print(f"Generated template workbook successfully, size: {len(content)} bytes")
